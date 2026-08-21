from datetime import date
from decimal import Decimal
from io import BytesIO

import uuid

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from cosecheros.models import Cosecha, Cosechero
from cosecheros.services import calcular_resumenes_cosecha
from ventas.models import DetalleAvance, OperacionVenta, Venta
from ventas.services import obtener_detalles_venta

from .models import Avance
from .excel import TEMPLATE_ROWS, generar_plantilla_avances
from .services import import_confirmed_rows, parse_fecha, parse_file_for_preview


class PlantillaAvancesTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.usuario = get_user_model().objects.create_user('plantilla', password='prueba')
        cls.activo = Cosechero.objects.create(
            nombre='Ana', apellido='Duplicada', cedula='', numero_cuenta_banco='12345',
            direccion='Local', telefono='', terreno_sembrado=Decimal('10.00'),
        )
        cls.duplicado = Cosechero.objects.create(
            nombre='Ana', apellido='Duplicada', cedula='', numero_cuenta_banco='67890',
            direccion='Local', telefono='', terreno_sembrado=Decimal('11.00'),
        )
        cls.inactivo = Cosechero.objects.create(
            nombre='Fuera', apellido='Catalogo', cedula='', numero_cuenta_banco=None,
            direccion='Local', telefono='', terreno_sembrado=Decimal('1.00'),
            is_active=False,
        )

    def test_descarga_requiere_autenticacion(self):
        response = self.client.get(reverse('avances_plantilla'))

        self.assertEqual(response.status_code, 302)

    def test_descarga_contiene_catalogo_formato_y_validaciones(self):
        self.client.force_login(self.usuario)
        response = self.client.get(reverse('avances_plantilla'))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertIn('plantilla_avances_', response['Content-Disposition'])

        workbook = load_workbook(BytesIO(b''.join(response.streaming_content)), data_only=False)
        self.assertEqual(workbook.sheetnames, ['Avances', 'Catalogo'])
        sheet = workbook['Avances']
        catalog = workbook['Catalogo']
        self.assertEqual(catalog.sheet_state, 'hidden')
        self.assertEqual(
            [sheet.cell(1, column).value for column in range(1, 8)],
            ['Tipo', 'Numero', 'Fecha', 'Cosechero', 'ID', 'Monto', 'Descripcion'],
        )
        self.assertEqual(sheet.freeze_panes, 'A2')
        self.assertEqual(sheet['C2'].number_format, '@')
        self.assertEqual(sheet['F2'].number_format, '"RD$" #,##0.00')
        self.assertIn('VLOOKUP(D2', sheet['E2'].value)
        self.assertTrue(sheet.protection.sheet)
        self.assertTrue(sheet['E2'].protection.locked)
        self.assertFalse(sheet['D2'].protection.locked)
        self.assertEqual(sheet.max_row, TEMPLATE_ROWS + 1)
        self.assertEqual(len(sheet.data_validations.dataValidation), 4)
        date_validation = next(
            validation for validation in sheet.data_validations.dataValidation
            if 'C2:C501' in str(validation.sqref)
        )
        self.assertEqual(date_validation.type, 'custom')
        self.assertIn('TEXT(DATE(', date_validation.formula1)

        selectors = [catalog.cell(row, 1).value for row in range(2, catalog.max_row + 1)]
        self.assertIn(f'{self.activo.id} — Ana Duplicada', selectors)
        self.assertIn(f'{self.duplicado.id} — Ana Duplicada', selectors)
        self.assertNotIn(f'{self.inactivo.id} — Fuera Catalogo', selectors)

    def test_plantilla_se_parsea_sin_resultado_calculado_de_formula(self):
        archivo = generar_plantilla_avances()
        workbook = load_workbook(archivo, data_only=False)
        sheet = workbook['Avances']
        sheet['A2'] = 'Efectivo'
        sheet['B2'] = 'EF-1'
        sheet['C2'] = '08-11-2026'
        sheet['D2'] = f'{self.activo.id} — Ana Duplicada'
        sheet['F2'] = Decimal('125.50')
        salida = BytesIO()
        workbook.save(salida)
        salida.name = 'plantilla_avances.xlsx'
        salida.seek(0)

        preview = parse_file_for_preview(salida)

        self.assertEqual(preview['formato'], 'unificado')
        self.assertEqual(preview['total_rows'], 1)
        self.assertEqual(preview['rows'][0]['cosechero_id'], self.activo.id)
        self.assertEqual(preview['rows'][0]['fecha'], '2026-11-08')
        self.assertEqual(Decimal(preview['rows'][0]['monto']), Decimal('125.50'))

    def test_fecha_textual_prioriza_dia_mes(self):
        self.assertEqual(parse_fecha('04-08-2026'), date(2026, 8, 4))
        self.assertEqual(parse_fecha('04/08/2026'), date(2026, 8, 4))
        self.assertIsNone(parse_fecha('31-02-2026'))


class ImportacionAvancesCosechaTests(TestCase):
    def test_no_reutiliza_venta_de_otra_cosecha(self):
        cosechero = Cosechero.objects.create(
            nombre='Importado', apellido='Prueba', cedula='', numero_cuenta_banco=None,
            direccion='Local', telefono='', terreno_sembrado=Decimal('1.00'),
        )
        anterior = Cosecha.objects.create(
            nombre='Anterior', fecha_inicio=date(2025, 1, 1), fecha_fin=date(2025, 12, 31),
        )
        actual = Cosecha.objects.create(
            nombre='Actual', fecha_inicio=date(2026, 1, 1), fecha_fin=date(2026, 12, 31),
        )
        Venta.objects.create(
            cosechero=cosechero, cosecha=anterior,
            fecha_venta=date(2026, 8, 8), total=Decimal('10.00'),
        )
        filas = [{
            'index': 1, 'cosechero_id': cosechero.id, 'monto': '25.00',
            'fecha': '2026-08-04', 'tipo_avance': 'efectivo',
            'numero': 'IMP-1', 'descripcion': 'Importación segura',
        }]

        resultado = import_confirmed_rows(filas, actual.id, 'Avance importado')

        self.assertEqual(resultado['errores'], [])
        self.assertEqual(Venta.objects.filter(cosecha=anterior).count(), 1)
        venta_actual = Venta.objects.get(cosecha=actual)
        self.assertEqual(venta_actual.cosechero, cosechero)
        self.assertEqual(venta_actual.fecha_venta, date(2026, 8, 8))
        self.assertEqual(venta_actual.detalle_avances.get().monto, Decimal('25.00'))

        venta_actual.impreso = True
        venta_actual.save(update_fields=['impreso'])
        filas[0]['numero'] = 'IMP-2'
        filas[0]['monto'] = '5.00'
        segundo = import_confirmed_rows(filas, actual.id, 'Avance importado')
        venta_actual.refresh_from_db()
        self.assertEqual(segundo['ventas_actualizadas'], 1)
        self.assertFalse(venta_actual.impreso)
        self.assertEqual(venta_actual.detalle_avances.count(), 2)


class CrudAvancesTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.usuario = get_user_model().objects.create_user('avances', password='prueba')
        cls.cosechero = Cosechero.objects.create(
            nombre='Ana', apellido='Avance', cedula='', numero_cuenta_banco=None,
            direccion='Local', telefono='', terreno_sembrado=Decimal('10.00'),
        )
        cls.otro_cosechero = Cosechero.objects.create(
            nombre='Bruno', apellido='Destino', cedula='', numero_cuenta_banco=None,
            direccion='Local', telefono='', terreno_sembrado=Decimal('12.00'),
        )
        cls.anterior = Cosecha.objects.create(
            nombre='Cosecha anterior CRUD',
            fecha_inicio=date(2025, 1, 1), fecha_fin=date(2025, 12, 31),
        )
        cls.actual = Cosecha.objects.create(
            nombre='Cosecha actual CRUD',
            fecha_inicio=date(2026, 1, 1), fecha_fin=date(2026, 12, 31),
        )

    def setUp(self):
        self.client.force_login(self.usuario)

    def crear_vinculado(self, *, monto='100.00', fecha=date(2026, 8, 4), estado='realizado'):
        venta = Venta.objects.create(
            cosechero=self.cosechero,
            cosecha=self.actual,
            fecha_venta=date(2026, 8, 8),
            total=Decimal(monto),
            impreso=True,
        )
        avance = Avance.objects.create(
            cosechero=self.cosechero,
            monto_pagado=Decimal(monto),
            fecha=fecha,
            descripcion='Prueba CRUD',
            tipo_avance='cheque',
            numero='CRUD-1',
            estado=estado,
        )
        DetalleAvance.objects.create(venta=venta, avance=avance, monto=avance.monto_pagado)
        return avance, venta

    def payload(self, **cambios):
        data = {
            'cosecha': str(self.actual.id),
            'cosechero': str(self.cosechero.id),
            'fecha': '2026-08-04',
            'tipo_avance': 'cheque',
            'numero': 'CRUD-WEB',
            'monto_pagado': '125.50',
            'descripcion': 'Creado desde CRUD',
            'estado': 'realizado',
            'idempotency_key': str(uuid.uuid4()),
        }
        data.update(cambios)
        return data

    def test_lista_abre_en_cosecha_actual_y_excluye_inactivos(self):
        activo, _ = self.crear_vinculado()
        inactivo = Avance.objects.create(
            cosechero=self.cosechero, monto_pagado=Decimal('50.00'),
            fecha=date(2025, 6, 1), tipo_avance='efectivo', estado='realizado',
            is_active=False,
        )

        response = self.client.get(reverse('avances'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, reverse('avance_detalle', args=[activo.id]))
        self.assertNotContains(response, f'/avances/{inactivo.id}/')
        self.assertEqual(response.context['cosecha_seleccionada'], self.actual)

    def test_lista_pagina_cincuenta_avances(self):
        Avance.objects.bulk_create([
            Avance(
                cosechero=self.cosechero,
                monto_pagado=Decimal('1.00'),
                fecha=date(2026, 8, 1),
                tipo_avance='efectivo',
                estado='realizado',
            )
            for _ in range(55)
        ])

        response = self.client.get(reverse('avances'), {'sin_cosecha': '1'})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.context['avances']), 50)
        self.assertEqual(response.context['page_obj'].paginator.count, 55)

    def test_filtros_y_resumen_de_tabla_son_consistentes(self):
        avance, _ = self.crear_vinculado(monto='123.45')
        Avance.objects.create(
            cosechero=self.cosechero,
            monto_pagado=Decimal('900.00'),
            fecha=date(2026, 8, 4),
            tipo_avance='deposito',
            estado='realizado',
        )

        response = self.client.get(reverse('avances'), {
            'cosecha': str(self.actual.id),
            'q': 'CRUD-1',
            'tipo': 'cheque',
            'estado': 'realizado',
            'desde': '2026-08-01',
            'hasta': '2026-08-08',
        })

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['resumen']['cantidad'], 1)
        self.assertEqual(response.context['resumen']['total'], Decimal('123.45'))
        self.assertEqual(list(response.context['avances']), [avance])

    def test_pantallas_crud_e_importador_renderizan(self):
        avance, _ = self.crear_vinculado()
        huerfano = Avance.objects.create(
            cosechero=self.cosechero,
            monto_pagado=Decimal('10.00'),
            fecha=date(2026, 8, 2),
            tipo_avance='efectivo',
            estado='realizado',
        )

        urls = [
            reverse('avance_nuevo'),
            reverse('avance_detalle', args=[avance.id]),
            reverse('avance_editar', args=[avance.id]),
            reverse('avance_vincular', args=[huerfano.id]),
            reverse('avances_importar'),
        ]
        for url in urls:
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 200)

    def test_creacion_individual_es_idempotente_y_ligada_a_ticket(self):
        clave = uuid.uuid4()
        payload = self.payload(idempotency_key=str(clave), numero='')

        primero = self.client.post(reverse('avance_nuevo'), payload)
        segundo = self.client.post(reverse('avance_nuevo'), payload)

        self.assertEqual(primero.status_code, 302)
        self.assertEqual(segundo.status_code, 302)
        self.assertEqual(Avance.objects.count(), 1)
        self.assertEqual(DetalleAvance.objects.count(), 1)
        self.assertEqual(OperacionVenta.objects.count(), 1)
        avance = Avance.objects.get()
        detalle = DetalleAvance.objects.get()
        self.assertEqual(avance.monto_pagado, Decimal('125.50'))
        self.assertEqual(detalle.monto, Decimal('125.50'))
        self.assertEqual(detalle.venta.fecha_venta, date(2026, 8, 8))

    def test_editar_monto_fecha_cosechero_y_cosecha_mueve_el_cargo(self):
        avance, venta_origen = self.crear_vinculado()
        data = self.payload(
            cosecha=str(self.anterior.id),
            cosechero=str(self.otro_cosechero.id),
            fecha='2025-09-10',
            monto_pagado='275.25',
            numero='CORREGIDO',
        )

        response = self.client.post(reverse('avance_editar', args=[avance.id]), data)

        self.assertEqual(response.status_code, 302)
        avance.refresh_from_db()
        detalle = DetalleAvance.objects.select_related('venta').get(avance=avance)
        venta_origen.refresh_from_db()
        self.assertEqual(avance.cosechero, self.otro_cosechero)
        self.assertEqual(avance.fecha, date(2025, 9, 10))
        self.assertEqual(avance.monto_pagado, Decimal('275.25'))
        self.assertEqual(detalle.monto, Decimal('275.25'))
        self.assertEqual(detalle.venta.cosecha, self.anterior)
        self.assertEqual(detalle.venta.cosechero, self.otro_cosechero)
        self.assertEqual(detalle.venta.fecha_venta, date(2025, 9, 13))
        self.assertEqual(venta_origen.total, Decimal('0.00'))
        self.assertFalse(venta_origen.impreso)

    def test_editar_fecha_dentro_de_la_semana_conserva_ticket(self):
        avance, venta = self.crear_vinculado()

        response = self.client.post(
            reverse('avance_editar', args=[avance.id]),
            self.payload(fecha='2026-08-07', monto_pagado='110.00'),
        )

        self.assertEqual(response.status_code, 302)
        detalle = DetalleAvance.objects.get(avance=avance)
        venta.refresh_from_db()
        self.assertEqual(detalle.venta_id, venta.id)
        self.assertEqual(Venta.objects.count(), 1)
        self.assertEqual(venta.total, Decimal('110.00'))

    def test_desactivar_excluye_y_restaurar_reincorpora(self):
        avance, venta = self.crear_vinculado(monto='80.00', estado='nulo')
        venta.refresh_from_db()
        self.assertEqual(venta.total, Decimal('80.00'))

        response = self.client.post(reverse('avance_desactivar', args=[avance.id]))
        self.assertEqual(response.status_code, 302)
        avance.refresh_from_db()
        venta.refresh_from_db()
        self.assertFalse(avance.is_active)
        self.assertEqual(venta.total, Decimal('0.00'))
        resumen = calcular_resumenes_cosecha(self.actual.id, [self.cosechero.id])[0]
        self.assertEqual(resumen['gastos_avances'], Decimal('0'))
        self.assertEqual(obtener_detalles_venta(venta)['avances'], [])

        response = self.client.post(reverse('avance_restaurar', args=[avance.id]))
        self.assertEqual(response.status_code, 302)
        avance.refresh_from_db()
        venta.refresh_from_db()
        self.assertTrue(avance.is_active)
        self.assertEqual(venta.total, Decimal('80.00'))
        self.assertEqual(avance.estado, 'nulo')
        resumen = calcular_resumenes_cosecha(self.actual.id, [self.cosechero.id])[0]
        self.assertEqual(resumen['gastos_avances'], Decimal('80.00'))

    def test_vincular_huerfano_lo_incorpora_una_sola_vez(self):
        avance = Avance.objects.create(
            cosechero=self.cosechero,
            monto_pagado=Decimal('45.75'),
            fecha=date(2026, 8, 5),
            tipo_avance='deposito',
            estado='realizado',
        )
        data = {
            'cosecha': str(self.actual.id),
            'cosechero': str(self.otro_cosechero.id),
            'fecha': '2026-08-05',
        }

        primero = self.client.post(reverse('avance_vincular', args=[avance.id]), data)
        segundo = self.client.post(reverse('avance_vincular', args=[avance.id]), data)

        self.assertEqual(primero.status_code, 302)
        self.assertEqual(segundo.status_code, 302)
        self.assertEqual(DetalleAvance.objects.filter(avance=avance).count(), 1)
        detalle = DetalleAvance.objects.select_related('venta').get(avance=avance)
        self.assertEqual(detalle.venta.cosechero, self.otro_cosechero)
        self.assertEqual(detalle.venta.cosecha, self.actual)
        self.assertEqual(detalle.venta.total, Decimal('45.75'))

    def test_get_no_desactiva_ni_restaura(self):
        avance, _ = self.crear_vinculado()
        self.assertEqual(
            self.client.get(reverse('avance_desactivar', args=[avance.id])).status_code,
            405,
        )
        avance.refresh_from_db()
        self.assertTrue(avance.is_active)
