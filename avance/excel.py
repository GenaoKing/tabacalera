"""Generación de la plantilla Excel para importar avances."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import quote_sheetname
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from cosecheros.models import Cosechero


TEMPLATE_ROWS = 500
HEADERS = ['Tipo', 'Numero', 'Fecha', 'Cosechero', 'ID', 'Monto', 'Descripcion']


def generar_plantilla_avances() -> BytesIO:
    """Construye una plantilla XLSX con el catálogo vigente de cosecheros."""
    cosecheros = list(
        Cosechero.objects.filter(is_active=True).order_by('nombre', 'apellido', 'id')
    )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Avances'
    catalog = workbook.create_sheet('Catalogo')

    header_fill = PatternFill('solid', fgColor='365314')
    header_font = Font(color='FFFFFF', bold=True)
    input_fill = PatternFill('solid', fgColor='F7FEE7')
    formula_fill = PatternFill('solid', fgColor='E2E8F0')

    for column, value in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=1, column=column, value=value)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')

    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = f'A1:G{TEMPLATE_ROWS + 1}'
    widths = {'A': 15, 'B': 18, 'C': 15, 'D': 38, 'E': 10, 'F': 18, 'G': 42}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    # Instrucciones visibles fuera de la tabla para no alterar sus encabezados.
    sheet['I1'] = 'Instrucciones'
    sheet['I1'].font = Font(bold=True, color='FFFFFF')
    sheet['I1'].fill = header_fill
    sheet['I2'] = 'Introduzca la fecha como texto: día-mes-año.'
    sheet['I3'] = 'Ejemplo inequívoco: 31-08-2026 (no use /).'
    sheet['I4'] = 'Seleccione el tipo y el cosechero en sus listas.'
    sheet['I5'] = 'La columna ID se completa automáticamente.'
    sheet.column_dimensions['I'].width = 54

    catalog.append(['Selector', 'ID', 'Nombre', 'Cuenta bancaria'])
    for cell in catalog[1]:
        cell.fill = header_fill
        cell.font = header_font

    for cosechero in cosecheros:
        nombre = f'{cosechero.nombre} {cosechero.apellido}'.strip()
        catalog.append([
            f'{cosechero.id} — {nombre}', cosechero.id, nombre,
            cosechero.numero_cuenta_banco or '',
        ])

    catalog_last_row = max(2, len(cosecheros) + 1)
    catalog_range = f'{quote_sheetname(catalog.title)}!$A$2:$A${catalog_last_row}'
    workbook.defined_names.add(DefinedName('ListaCosecheros', attr_text=catalog_range))

    type_validation = DataValidation(
        type='list', formula1='"Cheque,Deposito,Efectivo"', allow_blank=True,
    )
    type_validation.promptTitle = 'Tipo de avance'
    type_validation.prompt = 'Seleccione Cheque, Deposito o Efectivo.'
    type_validation.errorTitle = 'Tipo no válido'
    type_validation.error = 'Seleccione un tipo de la lista.'
    type_validation.errorStyle = 'stop'
    type_validation.showErrorMessage = True

    grower_validation = DataValidation(
        type='list', formula1='=ListaCosecheros', allow_blank=True,
    )
    grower_validation.promptTitle = 'Cosechero'
    grower_validation.prompt = 'Seleccione un cosechero activo.'
    grower_validation.errorTitle = 'Cosechero no válido'
    grower_validation.error = 'Seleccione un cosechero de la lista.'
    grower_validation.errorStyle = 'stop'
    grower_validation.showErrorMessage = True

    # La fecha se captura como texto para impedir que Excel aplique el orden
    # regional mm-dd al escribirla y luego la muestre engañosamente como dd-mm.
    # La fórmula valida estructura y fechas imposibles (incluidos bisiestos).
    date_validation = DataValidation(
        type='custom',
        formula1=(
            '=OR(C2="",IFERROR(AND(LEN(C2)=10,MID(C2,3,1)="-",'
            'MID(C2,6,1)="-",TEXT(DATE(VALUE(RIGHT(C2,4)),'
            'VALUE(MID(C2,4,2)),VALUE(LEFT(C2,2))),"dd-mm-yyyy")=C2),FALSE))'
        ),
        allow_blank=True,
    )
    date_validation.promptTitle = 'Fecha: día-mes-año'
    date_validation.prompt = 'Use dd-mm-aaaa. Ejemplo: 31-08-2026.'
    date_validation.errorTitle = 'Fecha no válida'
    date_validation.error = 'Escriba una fecha válida como 31-08-2026.'
    date_validation.errorStyle = 'stop'
    date_validation.showErrorMessage = True
    date_validation.showInputMessage = True

    amount_validation = DataValidation(
        type='decimal', operator='greaterThan', formula1='0', allow_blank=True,
    )
    amount_validation.errorTitle = 'Monto no válido'
    amount_validation.error = 'El monto debe ser mayor que cero.'
    amount_validation.errorStyle = 'stop'
    amount_validation.showErrorMessage = True

    for validation in (type_validation, grower_validation, date_validation, amount_validation):
        sheet.add_data_validation(validation)
    type_validation.add(f'A2:A{TEMPLATE_ROWS + 1}')
    grower_validation.add(f'D2:D{TEMPLATE_ROWS + 1}')
    date_validation.add(f'C2:C{TEMPLATE_ROWS + 1}')
    amount_validation.add(f'F2:F{TEMPLATE_ROWS + 1}')

    catalog_table = f"{quote_sheetname(catalog.title)}!$A$2:$B${catalog_last_row}"
    for row in range(2, TEMPLATE_ROWS + 2):
        for column in (1, 2, 3, 4, 6, 7):
            cell = sheet.cell(row=row, column=column)
            cell.protection = Protection(locked=False)
            cell.fill = input_fill

        # ``@`` obliga a conservar literalmente 08-11-2026; el servidor lo
        # convertirá a una fecha real como 8 de noviembre al importar.
        sheet.cell(row=row, column=3).number_format = '@'
        id_cell = sheet.cell(row=row, column=5)
        id_cell.value = f'=IFERROR(VLOOKUP(D{row},{catalog_table},2,FALSE),"")'
        id_cell.fill = formula_fill
        id_cell.protection = Protection(locked=True)
        # Las comillas hacen que ``RD$`` sea texto literal. Sin ellas,
        # openpyxl puede interpretar la D como parte de un formato de fecha.
        sheet.cell(row=row, column=6).number_format = '"RD$" #,##0.00'

    sheet['C1'].comment = Comment(
        'Formato obligatorio: dd-mm-aaaa. Ejemplo: 31-08-2026.', 'Tabacalera Genao',
    )
    sheet['E1'].comment = Comment(
        'Esta columna se calcula al seleccionar el cosechero.', 'Tabacalera Genao',
    )

    sheet.protection.sheet = True
    sheet.protection.autoFilter = False
    sheet.protection.sort = False
    sheet.protection.set_password('tabacalera')
    catalog.sheet_state = 'hidden'
    workbook.active = 0

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
